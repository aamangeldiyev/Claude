"""
Streaming Excel writer — appends detection rows as they arrive and
periodically flushes to disk to survive crashes mid-scan.

Falls back to streaming CSV if openpyxl is missing.
"""

import csv
import sys
from datetime import datetime
from pathlib import Path

HEADERS = [
    "Файл", "Папка", "Тип файла", "Страница", "Стр. всего",
    "Штамп", "Метод", "Уверенность, %", "Область (x1,y1,x2,y2)",
]
WIDTHS = [42, 60, 10, 10, 10, 30, 10, 14, 28]
FLUSH_EVERY = 1000  # rows


class StreamingExcelWriter:
    def __init__(self, output_path: Path | str):
        self.path = Path(output_path)
        self.row_count = 0
        self._unique_files: set[str] = set()
        self._engine = "excel"

        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            self._engine = "csv"
            self.path = self.path.with_suffix(".csv")
            self._csv_file = open(self.path, "w", newline="", encoding="utf-8-sig")
            self._csv_writer = csv.writer(self._csv_file, delimiter=";")
            self._csv_writer.writerow(HEADERS)
            print(f"[INFO] openpyxl missing — writing CSV: {self.path}", file=sys.stderr)
            return

        self._wb = openpyxl.Workbook()
        self._ws = self._wb.active
        self._ws.title = "Найденные штампы"

        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        hdr_font = Font(color="FFFFFF", bold=True, size=10)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, (hdr, width) in enumerate(zip(HEADERS, WIDTHS), 1):
            cell = self._ws.cell(row=1, column=col, value=hdr)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = center
            self._ws.column_dimensions[cell.column_letter].width = width
        self._ws.freeze_panes = "A2"
        self._ws.row_dimensions[1].height = 28

        self._next_row = 2
        # Eager save so the file exists even if scan crashes immediately.
        self._wb.save(self.path)

    def add_detection(self, result: dict) -> None:
        if self._engine == "csv":
            self._add_csv(result)
        else:
            self._add_xlsx(result)
        self._unique_files.add(result["file"])

    def _add_csv(self, result: dict) -> None:
        for det in result["detections"]:
            self._csv_writer.writerow([
                result["file_name"],
                result["folder"],
                result["file_type"],
                result["page"],
                result.get("total_pages", 1),
                det.get("template_name", "stamp"),
                det.get("method", "—"),
                f"{det['confidence']:.1%}",
                "{},{},{},{}".format(*det["bbox"]),
            ])
            self.row_count += 1
        if self.row_count % FLUSH_EVERY == 0:
            self._csv_file.flush()

    def _add_xlsx(self, result: dict) -> None:
        ws = self._ws
        for det in result["detections"]:
            row = [
                result["file_name"],
                result["folder"],
                result["file_type"],
                result["page"],
                result.get("total_pages", 1),
                det.get("template_name", "stamp"),
                det.get("method", "—"),
                round(det["confidence"] * 100, 1),
                "{},{},{},{}".format(*det["bbox"]),
            ]
            for col, val in enumerate(row, 1):
                ws.cell(row=self._next_row, column=col, value=val)
            self._next_row += 1
            self.row_count += 1
        if self.row_count % FLUSH_EVERY == 0:
            try:
                self._wb.save(self.path)
            except Exception as exc:
                print(f"[WARN] Periodic save failed: {exc}", file=sys.stderr)

    def close(self) -> None:
        if self._engine == "csv":
            try:
                self._csv_file.flush()
                self._csv_file.close()
            except Exception:
                pass
            return

        try:
            from openpyxl.styles import Font
            ws2 = self._wb.create_sheet("Сводка")
            ws2.column_dimensions["A"].width = 38
            ws2.column_dimensions["B"].width = 22
            now = datetime.now().strftime("%d.%m.%Y %H:%M")
            ws2["A1"] = "Отчёт о поиске штампов"
            ws2["A1"].font = Font(bold=True, size=13)
            for r, (label, value) in enumerate([
                ("Дата формирования:", now),
                ("Файлов со штампами:", len(self._unique_files)),
                ("Всего вхождений:", self.row_count),
            ], start=3):
                ws2.cell(row=r, column=1, value=label).font = Font(bold=True)
                ws2.cell(row=r, column=2, value=value)

            self._wb.save(self.path)
        except Exception as exc:
            print(f"[WARN] Final save failed: {exc}", file=sys.stderr)
